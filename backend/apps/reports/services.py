import csv
import io
from datetime import datetime, date, timedelta
from calendar import monthrange
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.db.models import Q, Sum, Count
from django.http import HttpResponse

from apps.accounts.models import Employee
from apps.attendance.models import Attendance
from apps.requests.models import ManualAttendanceRequest
from apps.leaves.models import Holiday, LeaveRequest
from apps.attendance.services import (
    get_dhaka_date,
    format_time_display,
    format_duration_display
)

class ReportService:
    @staticmethod
    def get_daily_report(report_date=None, department_id=None, employee_type=None, status_filter=None, type_filter=None):
        if report_date is None:
            report_date = get_dhaka_date()
        elif isinstance(report_date, str):
            report_date = datetime.strptime(report_date, "%Y-%m-%d").date()

        # Fetch active employees
        emp_qs = Employee.objects.select_related('department').filter(is_active=True)
        if department_id:
            emp_qs = emp_qs.filter(department_id=department_id)
        if employee_type:
            emp_qs = emp_qs.filter(employee_type=employee_type)

        employees = list(emp_qs.order_by('employee_id'))

        # Fetch attendances for date
        att_qs = (
            Attendance.objects
            .select_related('employee', 'approved_by')
            .filter(attendance_date=report_date)
        )
        att_map = {att.employee_id: att for att in att_qs}

        # Check for registered company holiday on this date
        holiday = Holiday.objects.filter(date=report_date).first()

        # Check for approved leaves covering this date
        approved_leaves = (
            LeaveRequest.objects
            .filter(status=LeaveRequest.Status.APPROVED, start_date__lte=report_date, end_date__gte=report_date)
            .select_related('employee')
        )
        leave_map = {leave.employee_id: leave for leave in approved_leaves}

        # Count pending manual requests for date
        pending_requests_count = ManualAttendanceRequest.objects.filter(
            attendance_date=report_date,
            status=ManualAttendanceRequest.Status.PENDING
        ).count()

        pending_approval_count = att_qs.filter(
            approval_status=Attendance.ApprovalStatus.PENDING
        ).count()
        rejected_count = att_qs.filter(
            approval_status=Attendance.ApprovalStatus.REJECTED
        ).count()

        rows = []
        present_count = 0
        late_count = 0
        absent_count = 0
        on_leave_count = 0
        holiday_count = 0
        checked_in_count = 0
        checked_out_count = 0
        incomplete_count = 0

        for emp in employees:
            att = att_map.get(emp.id)
            if att:
                # Only APPROVED days feed the official tally. Pending and
                # rejected rows are still listed (the admin needs to see the
                # work) but must not inflate present/late/minutes for payroll.
                is_approved = att.approval_status == Attendance.ApprovalStatus.APPROVED
                is_checked_in = bool(att.check_in_time)
                is_checked_out = bool(att.check_out_time)

                if is_approved:
                    if is_checked_in:
                        checked_in_count += 1
                    if is_checked_out:
                        checked_out_count += 1
                    else:
                        incomplete_count += 1

                    if att.status == Attendance.Status.PRESENT:
                        present_count += 1
                    elif att.status == Attendance.Status.LATE:
                        late_count += 1

                row = {
                    'employee_id': emp.employee_id,
                    'employee_name': emp.full_name,
                    'department': emp.department.name if emp.department else "--",
                    'employee_type': emp.employee_type,
                    'employee_type_display': emp.get_employee_type_display(),
                    'designation': emp.designation,
                    'check_in': format_time_display(att.check_in_time),
                    'check_out': format_time_display(att.check_out_time) if att.check_out_time else "--",
                    'duration': format_duration_display(att.working_duration_minutes),
                    'duration_minutes': att.working_duration_minutes or 0,
                    'type': att.attendance_type,
                    'status': att.status,
                    'approval_status': att.approval_status,
                    'approved_by': att.approved_by.username if att.approved_by else "--",
                    'remarks': att.admin_remarks or "",
                    # GPS capture metadata
                    'check_in_latitude': float(att.check_in_latitude) if att.check_in_latitude is not None else None,
                    'check_in_longitude': float(att.check_in_longitude) if att.check_in_longitude is not None else None,
                    'check_in_accuracy': att.check_in_accuracy,
                    'check_in_address': att.check_in_address or "",
                    'check_out_latitude': float(att.check_out_latitude) if att.check_out_latitude is not None else None,
                    'check_out_longitude': float(att.check_out_longitude) if att.check_out_longitude is not None else None,
                    'check_out_accuracy': att.check_out_accuracy,
                    'check_out_address': att.check_out_address or "",
                    'is_on_duty': is_checked_in and not is_checked_out,
                }
            else:
                # Determine if employee is on approved leave, official holiday, or absent
                if holiday:
                    holiday_count += 1
                    status_name = 'HOLIDAY'
                    remarks_text = f"Holiday: {holiday.name}"
                elif emp.id in leave_map:
                    on_leave_count += 1
                    l_req = leave_map[emp.id]
                    status_name = 'ON_LEAVE'
                    remarks_text = f"Leave: {l_req.get_leave_type_display()}"
                else:
                    absent_count += 1
                    status_name = Attendance.Status.ABSENT
                    remarks_text = ""

                row = {
                    'employee_id': emp.employee_id,
                    'employee_name': emp.full_name,
                    'department': emp.department.name if emp.department else "--",
                    'employee_type': emp.employee_type,
                    'employee_type_display': emp.get_employee_type_display(),
                    'designation': emp.designation,
                    'check_in': "--",
                    'check_out': "--",
                    'duration': "--",
                    'duration_minutes': 0,
                    'type': "--",
                    'status': status_name,
                    'approval_status': None,
                    'approved_by': "--",
                    'remarks': remarks_text,
                    'check_in_latitude': None,
                    'check_in_longitude': None,
                    'check_in_accuracy': None,
                    'check_in_address': "",
                    'check_out_latitude': None,
                    'check_out_longitude': None,
                    'check_out_accuracy': None,
                    'check_out_address': "",
                    'is_on_duty': False,
                }

            # Apply filters
            if status_filter and row['status'] != status_filter:
                continue
            if type_filter and row['type'] != type_filter:
                continue

            rows.append(row)

        summary = {
            'date': str(report_date),
            'total_assistants': len(employees),
            'present': present_count,
            'late': late_count,
            'on_leave': on_leave_count,
            'holiday': holiday_count,
            'absent': absent_count,
            'checked_in': checked_in_count,
            'checked_out': checked_out_count,
            'on_duty': incomplete_count,
            'incomplete': incomplete_count,
            'pending_manual_requests': pending_requests_count,
            'pending_approval': pending_approval_count,
            'rejected': rejected_count,
            'filtered_count': len(rows)
        }

        return {'summary': summary, 'records': rows}

    @staticmethod
    def get_monthly_report(year, month, department_id=None, employee_type=None):
        num_days = monthrange(year, month)[1]
        start_date = date(year, month, 1)
        end_date = date(year, month, num_days)

        emp_qs = Employee.objects.select_related('department').filter(is_active=True)
        if department_id:
            emp_qs = emp_qs.filter(department_id=department_id)
        if employee_type:
            emp_qs = emp_qs.filter(employee_type=employee_type)
        employees = list(emp_qs.order_by('employee_id'))

        # Fetch attendances for the month
        attendances = list(
            Attendance.objects
            .filter(
                attendance_date__gte=start_date,
                attendance_date__lte=end_date,
                # Payroll-grade totals only ever include admin-approved days.
                approval_status=Attendance.ApprovalStatus.APPROVED,
            )
            .values(
                'employee_id', 'attendance_date', 'status', 'attendance_type',
                'check_in_time', 'check_out_time', 'working_duration_minutes',
                'check_in_address', 'check_out_address', 'admin_remarks'
            )
        )
        att_map = {(att['employee_id'], att['attendance_date']): att for att in attendances}

        pending_approval_count = Attendance.objects.filter(
            attendance_date__gte=start_date,
            attendance_date__lte=end_date,
            approval_status=Attendance.ApprovalStatus.PENDING,
        ).count()

        # Days that exist but are NOT approved. They must be classified as
        # "Awaiting Approval"/"Rejected" rather than falling through to
        # ABSENT, which would misreport every unsigned day as a formal absence.
        unapproved_map = {
            (row['employee_id'], row['attendance_date']): row
            for row in Attendance.objects
            .filter(
                attendance_date__gte=start_date,
                attendance_date__lte=end_date,
            )
            .exclude(approval_status=Attendance.ApprovalStatus.APPROVED)
            .values('employee_id', 'attendance_date', 'approval_status', 'status')
        }

        # Fetch holidays in this month
        holidays_qs = Holiday.objects.filter(date__gte=start_date, date__lte=end_date)
        holiday_map = {h.date: h.name for h in holidays_qs}

        # Fetch approved leaves in this month
        approved_leaves = list(
            LeaveRequest.objects
            .filter(status=LeaveRequest.Status.APPROVED, start_date__lte=end_date, end_date__gte=start_date)
            .values('employee_id', 'leave_type', 'start_date', 'end_date', 'reason')
        )
        leaves_by_emp = {}
        for l in approved_leaves:
            leaves_by_emp.setdefault(l['employee_id'], []).append(l)

        # Build days metadata for the month
        days_metadata = []
        for d in range(1, num_days + 1):
            cur_date = date(year, month, d)
            is_friday = (cur_date.weekday() == 4)
            is_h = cur_date in holiday_map
            days_metadata.append({
                'day': d,
                'date': cur_date.isoformat(),
                'weekday': cur_date.strftime("%a"),
                'is_weekend': is_friday,
                'is_holiday': is_h,
                'holiday_name': holiday_map.get(cur_date, "")
            })

        today = get_dhaka_date()
        records = []

        for emp in employees:
            e_leaves = leaves_by_emp.get(emp.id, [])
            emp_days = []

            present_count = 0
            late_count = 0
            manual_count = 0
            leave_count = 0
            holiday_count = 0
            weekend_count = 0
            absent_count = 0
            total_minutes = 0
            pending_count = 0
            rejected_count = 0

            for d in range(1, num_days + 1):
                cur_date = date(year, month, d)
                is_friday = (cur_date.weekday() == 4)
                att = att_map.get((emp.id, cur_date))

                if att:
                    st = att['status']
                    a_type = att['attendance_type']
                    mins = att['working_duration_minutes'] or 0
                    total_minutes += mins

                    if st == Attendance.Status.PRESENT:
                        present_count += 1
                    elif st == Attendance.Status.LATE:
                        late_count += 1

                    if a_type == Attendance.Type.MANUAL:
                        manual_count += 1
                        code = 'M'
                        label = 'Manual (Approved)'
                    elif st == Attendance.Status.LATE:
                        code = 'L'
                        label = 'Late Arrival'
                    else:
                        code = 'P'
                        label = 'Present'

                    emp_days.append({
                        'day': d,
                        'date': cur_date.isoformat(),
                        'weekday': cur_date.strftime("%a"),
                        'status': st,
                        'code': code,
                        'label': label,
                        'check_in': format_time_display(att['check_in_time']),
                        'check_out': format_time_display(att['check_out_time']) if att['check_out_time'] else "--",
                        'duration': format_duration_display(mins),
                        'duration_minutes': mins,
                        'type': a_type,
                        'check_in_address': att.get('check_in_address') or "",
                        'check_out_address': att.get('check_out_address') or "",
                        'remarks': att.get('admin_remarks') or ""
                    })

                elif cur_date in holiday_map:
                    holiday_count += 1
                    emp_days.append({
                        'day': d,
                        'date': cur_date.isoformat(),
                        'weekday': cur_date.strftime("%a"),
                        'status': 'HOLIDAY',
                        'code': 'H',
                        'label': holiday_map[cur_date],
                        'check_in': "--",
                        'check_out': "--",
                        'duration': "--",
                        'duration_minutes': 0,
                        'type': "--",
                        'check_in_address': "",
                        'check_out_address': "",
                        'remarks': f"Official Holiday: {holiday_map[cur_date]}"
                    })

                else:
                    # Check approved leave
                    matched_leave = next(
                        (l for l in e_leaves if l['start_date'] <= cur_date <= l['end_date']),
                        None
                    )
                    if matched_leave:
                        leave_count += 1
                        leave_type_display = matched_leave['leave_type'].replace('_', ' ').title()
                        emp_days.append({
                            'day': d,
                            'date': cur_date.isoformat(),
                            'weekday': cur_date.strftime("%a"),
                            'status': 'ON_LEAVE',
                            'code': 'LV',
                            'label': f"{leave_type_display} Leave",
                            'check_in': "--",
                            'check_out': "--",
                            'duration': "--",
                            'duration_minutes': 0,
                            'type': "--",
                            'check_in_address': "",
                            'check_out_address': "",
                            'remarks': matched_leave.get('reason') or f"Approved {leave_type_display} Leave"
                        })
                    elif is_friday:
                        weekend_count += 1
                        emp_days.append({
                            'day': d,
                            'date': cur_date.isoformat(),
                            'weekday': cur_date.strftime("%a"),
                            'status': 'WEEKEND',
                            'code': 'W',
                            'label': 'Weekly Off (Friday)',
                            'check_in': "--",
                            'check_out': "--",
                            'duration': "--",
                            'duration_minutes': 0,
                            'type': "--",
                            'check_in_address': "",
                            'check_out_address': "",
                            'remarks': "Weekly Holiday"
                        })
                    elif cur_date > today:
                        emp_days.append({
                            'day': d,
                            'date': cur_date.isoformat(),
                            'weekday': cur_date.strftime("%a"),
                            'status': 'FUTURE',
                            'code': '-',
                            'label': 'Upcoming Date',
                            'check_in': "--",
                            'check_out': "--",
                            'duration': "--",
                            'duration_minutes': 0,
                            'type': "--",
                            'check_in_address': "",
                            'check_out_address': "",
                            'remarks': ""
                        })
                    elif (emp.id, cur_date) in unapproved_map:
                        unapproved = unapproved_map[(emp.id, cur_date)]
                        is_rejected = unapproved['approval_status'] == Attendance.ApprovalStatus.REJECTED
                        if is_rejected:
                            rejected_count += 1
                        else:
                            pending_count += 1
                        emp_days.append({
                            'day': d,
                            'date': cur_date.isoformat(),
                            'weekday': cur_date.strftime("%a"),
                            'status': 'REJECTED' if is_rejected else 'PENDING_APPROVAL',
                            'code': 'R' if is_rejected else 'W',
                            'label': 'Rejected' if is_rejected else 'Awaiting Approval',
                            'check_in': "--",
                            'check_out': "--",
                            'duration': "--",
                            'duration_minutes': 0,
                            'type': unapproved['status'],
                            'check_in_address': "",
                            'check_out_address': "",
                            'remarks': (
                                "Rejected by admin - excluded from payroll; void to allow re-punch"
                                if is_rejected else
                                "Awaiting admin approval - not yet counted in payroll totals"
                            )
                        })
                    else:
                        absent_count += 1
                        emp_days.append({
                            'day': d,
                            'date': cur_date.isoformat(),
                            'weekday': cur_date.strftime("%a"),
                            'status': 'ABSENT',
                            'code': 'A',
                            'label': 'Absent',
                            'check_in': "--",
                            'check_out': "--",
                            'duration': "--",
                            'duration_minutes': 0,
                            'type': "--",
                            'check_in_address': "",
                            'check_out_address': "",
                            'remarks': "Unexcused absence / no mobile check-in recorded"
                        })

            records.append({
                'employee_id': emp.employee_id,
                'employee_name': emp.full_name,
                'department': emp.department.name if emp.department else "--",
                'employee_type': emp.employee_type,
                'employee_type_display': emp.get_employee_type_display(),
                'designation': emp.designation,
                'present_days': present_count,
                'late_days': late_count,
                'leave_days': leave_count,
                'holiday_days': holiday_count,
                'weekend_days': weekend_count,
                'absent_days': absent_count,
                'pending_approval_days': pending_count,
                'rejected_days': rejected_count,
                'manual_days': manual_count,
                'total_working_hours': format_duration_display(total_minutes),
                'total_working_minutes': total_minutes,
                'days': emp_days
            })

        return {
            'year': year,
            'month': month,
            'total_days': num_days,
            'days_metadata': days_metadata,
            'total_holidays': len(holiday_map),
            'total_employees': len(employees),
            'pending_approval_days': pending_approval_count,
            'records': records
        }

    @staticmethod
    def get_employee_report(employee_id_or_pk, start_date=None, end_date=None, employee_type=None):
        if isinstance(employee_id_or_pk, int) or str(employee_id_or_pk).isdigit():
            emp = Employee.objects.select_related('department', 'user').filter(id=int(employee_id_or_pk)).first()
        else:
            emp = Employee.objects.select_related('department', 'user').filter(employee_id=employee_id_or_pk).first()

        if not emp:
            return None

        # Employee-type scoping guard: an employee outside the requested type is
        # treated as "not found" so a type-filtered report can never expose
        # another group's dossier.
        if employee_type and str(emp.employee_type or '') != str(employee_type):
            return None

        today = get_dhaka_date()
        if not start_date:
            start_date = date(today.year, today.month, 1)
        elif isinstance(start_date, str):
            start_date = datetime.strptime(start_date, "%Y-%m-%d").date()

        if not end_date:
            end_date = today
        elif isinstance(end_date, str):
            end_date = datetime.strptime(end_date, "%Y-%m-%d").date()

        # Official holidays and approved leaves covering the window
        holiday_map = {
            h.date: h.name
            for h in Holiday.objects.filter(date__gte=start_date, date__lte=end_date)
        }
        approved_leaves = list(
            LeaveRequest.objects
            .filter(
                employee=emp,
                status=LeaveRequest.Status.APPROVED,
                start_date__lte=end_date,
                end_date__gte=start_date,
            )
            .values('leave_type', 'start_date', 'end_date', 'reason')
        )

        attendances = (
            Attendance.objects
            .select_related('approved_by')
            .filter(employee=emp, attendance_date__gte=start_date, attendance_date__lte=end_date)
        )
        att_by_date = {att.attendance_date: att for att in attendances}

        daily_records = []
        present_count = 0
        late_count = 0
        manual_count = 0
        total_minutes = 0
        pending_approval_count = 0
        rejected_count = 0
        leave_count = 0
        holiday_count = 0
        weekend_count = 0
        future_count = 0
        absent_count = 0

        num_days = (end_date - start_date).days + 1
        # Emit a row for EVERY calendar day of the window so the dossier shows
        # exactly which days were business absences, leave, holidays or weekends
        # instead of silently omitting them (and mislabelling holidays/off-days
        # as absences in the aggregate).
        for offset in range(num_days):
            cur_date = start_date + timedelta(days=offset)
            att = att_by_date.get(cur_date)

            if att:
                is_approved = att.approval_status == Attendance.ApprovalStatus.APPROVED
                if not is_approved:
                    if att.approval_status == Attendance.ApprovalStatus.PENDING:
                        pending_approval_count += 1
                    else:
                        rejected_count += 1
                    # Still listed for the reviewer's benefit, but contributes
                    # nothing to the payroll-grade totals below.
                    if att.approval_status == Attendance.ApprovalStatus.REJECTED:
                        status_name = 'REJECTED'
                        remarks_text = 'Rejected by admin - excluded from payroll'
                    else:
                        status_name = 'PENDING_APPROVAL'
                        remarks_text = 'Awaiting admin approval - not yet counted'
                    daily_records.append({
                        'date': str(cur_date),
                        'check_in': format_time_display(att.check_in_time),
                        'check_out': format_time_display(att.check_out_time) if att.check_out_time else "--",
                        'duration': '--',
                        'duration_minutes': 0,
                        'type': att.attendance_type,
                        'status': status_name,
                        'approval_status': att.approval_status,
                        'approved_by': att.approved_by.username if att.approved_by else "--",
                        'check_in_latitude': float(att.check_in_latitude) if att.check_in_latitude is not None else None,
                        'check_in_longitude': float(att.check_in_longitude) if att.check_in_longitude is not None else None,
                        'check_in_address': att.check_in_address or "",
                        'check_out_latitude': float(att.check_out_latitude) if att.check_out_latitude is not None else None,
                        'check_out_longitude': float(att.check_out_longitude) if att.check_out_longitude is not None else None,
                        'check_out_address': att.check_out_address or "",
                        'remarks': att.admin_remarks or remarks_text,
                    })
                    continue

                if att.status == Attendance.Status.PRESENT:
                    present_count += 1
                elif att.status == Attendance.Status.LATE:
                    late_count += 1

                if att.attendance_type == Attendance.Type.MANUAL:
                    manual_count += 1

                mins = att.working_duration_minutes or 0
                total_minutes += mins

                daily_records.append({
                    'date': str(cur_date),
                    'check_in': format_time_display(att.check_in_time),
                    'check_out': format_time_display(att.check_out_time) if att.check_out_time else "--",
                    'duration': format_duration_display(att.working_duration_minutes),
                    'type': att.attendance_type,
                    'status': att.status,
                    'approval_status': att.approval_status,
                    'approved_by': att.approved_by.username if att.approved_by else "--",
                    'remarks': att.admin_remarks or "",
                    'check_in_latitude': float(att.check_in_latitude) if att.check_in_latitude is not None else None,
                    'check_in_longitude': float(att.check_in_longitude) if att.check_in_longitude is not None else None,
                    'check_in_address': att.check_in_address or "",
                    'check_out_latitude': float(att.check_out_latitude) if att.check_out_latitude is not None else None,
                    'check_out_longitude': float(att.check_out_longitude) if att.check_out_longitude is not None else None,
                    'check_out_address': att.check_out_address or "",
                })
            elif cur_date in holiday_map:
                holiday_count += 1
                daily_records.append({
                    'date': str(cur_date),
                    'check_in': "--",
                    'check_out': "--",
                    'duration': "--",
                    'duration_minutes': 0,
                    'type': "--",
                    'status': 'HOLIDAY',
                    'approval_status': None,
                    'approved_by': "--",
                    'remarks': f"Official Holiday: {holiday_map[cur_date]}",
                })
            else:
                matched_leave = next(
                    (l for l in approved_leaves if l['start_date'] <= cur_date <= l['end_date']),
                    None
                )
                if matched_leave:
                    leave_count += 1
                    leave_type_display = matched_leave['leave_type'].replace('_', ' ').title()
                    daily_records.append({
                        'date': str(cur_date),
                        'check_in': "--",
                        'check_out': "--",
                        'duration': "--",
                        'duration_minutes': 0,
                        'type': "--",
                        'status': 'ON_LEAVE',
                        'approval_status': None,
                        'approved_by': "--",
                        'remarks': matched_leave.get('reason') or f"Approved {leave_type_display} Leave",
                    })
                elif cur_date.weekday() == 4:  # Friday = weekly off (Asia/Dhaka)
                    weekend_count += 1
                    daily_records.append({
                        'date': str(cur_date),
                        'check_in': "--",
                        'check_out': "--",
                        'duration': "--",
                        'duration_minutes': 0,
                        'type': "--",
                        'status': 'WEEKEND',
                        'approval_status': None,
                        'approved_by': "--",
                        'remarks': "Weekly Holiday (Friday)",
                    })
                elif cur_date > today:
                    future_count += 1
                    daily_records.append({
                        'date': str(cur_date),
                        'check_in': "--",
                        'check_out': "--",
                        'duration': "--",
                        'duration_minutes': 0,
                        'type': "--",
                        'status': 'FUTURE',
                        'approval_status': None,
                        'approved_by': "--",
                        'remarks': "Upcoming date",
                    })
                else:
                    absent_count += 1
                    daily_records.append({
                        'date': str(cur_date),
                        'check_in': "--",
                        'check_out': "--",
                        'duration': "--",
                        'duration_minutes': 0,
                        'type': "--",
                        'status': 'ABSENT',
                        'approval_status': None,
                        'approved_by': "--",
                        'remarks': "Unexcused absence / no mobile check-in recorded",
                    })

        return {
            'employee': {
                'id': emp.id,
                'employee_id': emp.employee_id,
                'full_name': emp.full_name,
                'phone': emp.phone,
                'email': emp.user.email,
                'department': emp.department.name if emp.department else "--",
                'employee_type': emp.employee_type,
                'employee_type_display': emp.get_employee_type_display(),
                'designation': emp.designation,
                'joining_date': str(emp.joining_date),
                'is_active': emp.is_active
            },
            'summary': {
                'start_date': str(start_date),
                'end_date': str(end_date),
                'total_days_range': num_days,
                'present_days': present_count,
                'late_days': late_count,
                'absent_days': absent_count,
                'leave_days': leave_count,
                'holiday_days': holiday_count,
                'weekend_days': weekend_count,
                'future_days': future_count,
                'manual_days': manual_count,
                'pending_approval_days': pending_approval_count,
                'rejected_days': rejected_count,
                'total_working_hours': format_duration_display(total_minutes),
                'total_working_minutes': total_minutes
            },
            'daily_records': daily_records
        }

    # ================= Export Engine =================
    # ================= Export Engine =================
    @staticmethod
    def export_csv_response(filename, headers, rows):
        import codecs
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        # Write UTF-8 BOM so MS Excel and Windows spreadsheet viewers open international text cleanly
        response.write(codecs.BOM_UTF8)
        writer = csv.writer(response)
        writer.writerow(headers)
        for r in rows:
            writer.writerow(r)
        return response

    @staticmethod
    def export_xlsx_response(filename, title, headers, rows, metadata=None, summary_stats=None, report_type='generic'):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Official Register"
        ws.views.sheetView[0].showGridLines = True

        # Typography and styling definitions
        title_font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
        subtitle_font = Font(name='Calibri', size=10.5, bold=True, color='FFFFFF')
        meta_label_font = Font(name='Calibri', size=9, bold=True, color='475569')
        meta_val_font = Font(name='Calibri', size=9, bold=False, color='0F172A')
        header_font = Font(name='Calibri', size=9.5, bold=True, color='FFFFFF')
        data_font = Font(name='Calibri', size=9)
        bold_data_font = Font(name='Calibri', size=9, bold=True)
        sig_title_font = Font(name='Calibri', size=9.5, bold=True, color='1E293B')
        sig_sub_font = Font(name='Calibri', size=8.5, italic=True, color='64748B')
        disclaimer_font = Font(name='Calibri', size=8, italic=True, color='64748B')

        navy_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid') # Slate 900
        blue_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid') # Blue 800
        header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # Slate 800
        zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid') # Slate 50
        meta_bg = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid') # Slate 100

        # Status badge colors
        status_styles = {
            'PRESENT': (
                PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='15803D')
            ),
            'P': (
                PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='15803D')
            ),
            'LATE': (
                PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='B45309')
            ),
            'L': (
                PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='B45309')
            ),
            'MANUAL': (
                PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='1D4ED8')
            ),
            'M': (
                PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='1D4ED8')
            ),
            'ON_LEAVE': (
                PatternFill(start_color='EDE9FE', end_color='EDE9FE', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='6D28D9')
            ),
            'LV': (
                PatternFill(start_color='EDE9FE', end_color='EDE9FE', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='6D28D9')
            ),
            'HOLIDAY': (
                PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='0369A1')
            ),
            'H': (
                PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='0369A1')
            ),
            'WEEKEND': (
                PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='64748B')
            ),
            'W': (
                PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='64748B')
            ),
            'ABSENT': (
                PatternFill(start_color='FFE4E6', end_color='FFE4E6', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='BE123C')
            ),
            'A': (
                PatternFill(start_color='FFE4E6', end_color='FFE4E6', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='BE123C')
            ),
            'INCOMPLETE': (
                PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='475569')
            ),
            'INC': (
                PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid'),
                Font(name='Calibri', size=9, bold=True, color='475569')
            ),
        }

        thin_side = Side(style='thin', color='CBD5E1')
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(
            left=thin_side, right=thin_side, top=thin_side,
            bottom=Side(style='double', color='1E293B')
        )

        current_row = 1

        # 1. Organization Letterhead Banner
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        top_cell = ws.cell(row=current_row, column=1, value="FRA - FIELD RESEARCH ASSISTANTS")
        top_cell.font = title_font
        top_cell.fill = navy_fill
        top_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[current_row].height = 28
        current_row += 1

        # 2. Report Title & Subtitle Banner
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        sub_cell = ws.cell(row=current_row, column=1, value=title.upper())
        sub_cell.font = subtitle_font
        sub_cell.fill = blue_fill
        sub_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[current_row].height = 22
        current_row += 1

        # 3. Metadata Section
        if metadata:
            current_row += 1
            mid_col = max(len(headers) // 2 + 1, 5)
            for idx, (lbl, val) in enumerate(metadata):
                is_right = (idx % 2 == 1)
                r = current_row + (idx // 2)
                col_lbl = mid_col if is_right else 1
                col_val = (mid_col + 1) if is_right else 2

                c_lbl = ws.cell(row=r, column=col_lbl, value=f"{lbl}:")
                c_lbl.font = meta_label_font
                c_lbl.alignment = Alignment(horizontal='left', vertical='center')

                c_val = ws.cell(row=r, column=col_val, value=str(val))
                c_val.font = meta_val_font
                c_val.alignment = Alignment(horizontal='left', vertical='center')
            
            current_row += ((len(metadata) + 1) // 2)

        # 4. Executive Summary KPI Badges
        if summary_stats:
            current_row += 1
            ws.cell(row=current_row, column=1, value="EXECUTIVE ATTENDANCE SUMMARY").font = Font(name='Calibri', size=9, bold=True, color='1E40AF')
            current_row += 1
            
            for s_idx, (s_label, s_val) in enumerate(summary_stats):
                col_pos = s_idx + 1
                if col_pos <= len(headers):
                    c = ws.cell(row=current_row, column=col_pos, value=f"{s_label}\n{s_val}")
                    c.font = Font(name='Calibri', size=8.5, bold=True, color='0F172A')
                    c.fill = meta_bg
                    c.border = thin_border
                    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.row_dimensions[current_row].height = 28
            current_row += 1

        current_row += 1 # Empty buffer line before table headers

        # 5. Table Column Headers
        table_header_row = current_row
        for col_num, h_text in enumerate(headers, 1):
            cell = ws.cell(row=table_header_row, column=col_num, value=h_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[table_header_row].height = 26
        current_row += 1

        # 6. Data Rows
        for r_idx, r_data in enumerate(rows):
            is_odd = (r_idx % 2 == 1)
            row_fill = zebra_fill if is_odd else PatternFill(fill_type=None)
            ws.row_dimensions[current_row].height = 20

            for col_idx, val in enumerate(r_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                cell.fill = row_fill

                val_str = str(val or '').strip()
                header_name = headers[col_idx - 1].lower() if col_idx - 1 < len(headers) else ''

                if 'sl' in header_name or col_idx == 1:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif 'id' in header_name:
                    cell.font = bold_data_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif 'name' in header_name:
                    cell.font = bold_data_font
                    cell.alignment = Alignment(horizontal='left', vertical='center')
                elif 'check in' in header_name or 'check out' in header_name or 'date' in header_name:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif 'duration' in header_name or 'hours' in header_name:
                    cell.font = bold_data_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif 'status' in header_name or val_str.upper() in status_styles:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    status_key = val_str.upper().replace(' ', '_')
                    if status_key in status_styles:
                        fill, font = status_styles[status_key]
                        cell.fill = fill
                        cell.font = font
                elif isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

            current_row += 1

        # 7. Total Row
        total_cell = ws.cell(row=current_row, column=1, value=f"Total Records: {len(rows)}")
        total_cell.font = bold_data_font
        for c in range(1, len(headers) + 1):
            ws.cell(row=current_row, column=c).border = double_bottom
        current_row += 3

        # 8. Formal 3-Tier HR Verification & Signature Block
        sig_col_1 = 2
        sig_col_2 = max(len(headers) // 2, 5)
        sig_col_3 = max(len(headers) - 2, sig_col_2 + 3)

        ws.cell(row=current_row, column=sig_col_1, value="___________________________________").font = sig_sub_font
        ws.cell(row=current_row, column=sig_col_2, value="___________________________________").font = sig_sub_font
        ws.cell(row=current_row, column=sig_col_3, value="___________________________________").font = sig_sub_font
        current_row += 1

        ws.cell(row=current_row, column=sig_col_1, value="Prepared By: HR Operations Assistant").font = sig_title_font
        ws.cell(row=current_row, column=sig_col_2, value="Verified By: Field Operations Lead").font = sig_title_font
        ws.cell(row=current_row, column=sig_col_3, value="Approved By: Head of HR & Admin").font = sig_title_font
        current_row += 1

        ws.cell(row=current_row, column=sig_col_1, value="Date: _____________________________").font = sig_sub_font
        ws.cell(row=current_row, column=sig_col_2, value="Date: _____________________________").font = sig_sub_font
        ws.cell(row=current_row, column=sig_col_3, value="Date: _____________________________").font = sig_sub_font
        current_row += 2

        # 9. Audit & Legal Disclaimer Footer
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        disc_cell = ws.cell(
            row=current_row,
            column=1,
            value="OFFICIAL DOCUMENT NOTICE: This is an authentic system-generated report from FRA - Field Research Assistants with authoritative server-side timestamps (Asia/Dhaka). Valid for HR payroll calculation, statutory compliance, and corporate audit."
        )
        disc_cell.font = disclaimer_font
        disc_cell.alignment = Alignment(horizontal='center', vertical='center')

        # Auto-adjust column widths
        for col in ws.columns:
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row >= table_header_row and cell.value and cell.coordinate not in ws.merged_cells:
                    lines = str(cell.value).split('\n')
                    max_len = max(max_len, max(len(line) for line in lines))
            ws.column_dimensions[col_letter].width = max(max_len + 2, 8)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
